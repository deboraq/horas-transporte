"""Genera Control_Horas.xlsx listo para subir a Google Sheets."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parent / "Control_Horas.xlsx"

PERSONAL = [
    (319, "MOREYRA, Leandro"),
    (324, "ALVAREZ, Mariano"),
    (2698, "RUIZ, Sergio"),
    (2794, "ACOSTA, Juan"),
    (247, "DIAZ, Silvia"),
]


def hdr(ws, row, cols, color="1F4E79"):
    fill = PatternFill("solid", fgColor=color)
    for i, t in enumerate(cols, 1):
        c = ws.cell(row=row, column=i, value=t)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(horizontal="center")


def main():
    wb = Workbook()

    # --- CARGA DIARIA ---
    ws = wb.active
    ws.title = "Carga"
    hdr(ws, 1, [
        "Fecha", "Legajo", "Nombre (auto)", "Entrada", "Salida",
        "Citacion", "Franco", "Feriado", "Turno",
    ])
    ws["A2"] = "26/05/2026"
    ws["B2"] = 319
    ws["D2"] = "12:56"
    ws["E2"] = "23:02"
    ws["F2"] = "13:10"
    ws["I2"] = "cronograma"
    ws["A3"] = "27/05/2026"
    ws["B3"] = 319
    ws["D3"] = "06:04"
    ws["E3"] = "19:28"
    ws["F3"] = "06:10"
    ws["I3"] = "cronograma"

    dv_franco = DataValidation(type="list", formula1='"si,no"', allow_blank=True)
    ws.add_data_validation(dv_franco)
    dv_franco.add("G2:G2000")
    dv_fer = DataValidation(type="list", formula1='"si,no"', allow_blank=True)
    ws.add_data_validation(dv_fer)
    dv_fer.add("H2:H2000")
    dv_turno = DataValidation(type="list", formula1='"cronograma,transporte"', allow_blank=True)
    ws.add_data_validation(dv_turno)
    dv_turno.add("I2:I2000")

    for col, w in zip("ABCDEFGHI", [12, 8, 28, 10, 10, 10, 8, 8, 14]):
        ws.column_dimensions[col].width = w

    # --- PERSONAL ---
    wp = wb.create_sheet("Personal")
    hdr(wp, 1, ["Legajo", "Apellido y Nombre"])
    for i, (leg, nom) in enumerate(PERSONAL, 2):
        wp.append([leg, nom])
    for row in range(2, 2002):
        ws.cell(row=row, column=3, value=f'=IFERROR(VLOOKUP(B{row},Personal!A:B,2,FALSE),"")')
    wp.column_dimensions["A"].width = 10
    wp.column_dimensions["B"].width = 32

    # --- RESULTADO (lo llena el script de Google o la web) ---
    wr = wb.create_sheet("Resultado")
    hdr(wr, 1, [
        "Fecha", "Legajo", "Nombre", "Total", "Normales", "50%", "100%",
        "Nocturnas", "Descanso", "Descanso 100%", "Cenas",
    ], color="375623")

    # --- INSTRUCCIONES ---
    wi = wb.create_sheet("Como usar")
    pasos = [
        "CONTROL DE HORAS — TRANSPORTE",
        "",
        "OPCION A — GOOGLE SHEETS (recomendado)",
        "1. Subí este archivo a Google Drive",
        "2. Abrilo con Google Hojas de cálculo",
        "3. Extensiones > Apps Script",
        "4. Pegá el contenido del archivo Code.gs (carpeta sheets)",
        "5. Guardá y recargá la hoja: menú Horas > Calcular período",
        "",
        "OPCION B — MAS SIMPLE (sin Google)",
        "1. Doble clic en: control_horas.html (en la carpeta Horas transporte)",
        "2. Cargá los horarios en el navegador",
        "3. Botón Calcular y Descargar Excel",
        "",
        "HOJA CARGA: solo completá Fecha, Legajo, Entrada, Salida, Citacion, Turno",
        "  Franco = si | Feriado = si | Turno = cronograma o transporte",
        "",
        "Período habitual: del 26 al 25 de cada mes.",
    ]
    for i, t in enumerate(pasos, 1):
        wi.cell(row=i, column=1, value=t)
    wi.column_dimensions["A"].width = 70

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
