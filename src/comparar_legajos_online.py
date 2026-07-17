"""
Compara cálculo propio vs exportación Legajos Online.
Uso: python src/comparar_legajos_online.py archivo.xlsx [--legajo 324]
"""
from __future__ import annotations

import argparse
import sys
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from calculo_horas import ResultadoDia, _fmt_min, _parse_date, calcular_dia
from importar_legajos_online import importar_legajos_online, legajo_desde_nombre

BASE = Path(__file__).resolve().parent.parent
PLANTILLA = BASE / "data" / "plantilla.xlsx"
SALIDA = BASE / "data" / "comparaciones"


def _leer_personal() -> dict[int, dict]:
    wb = load_workbook(PLANTILLA, data_only=True)
    ws = wb["Personal"]
    personal = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        personal[int(row[0])] = {"nombre": str(row[1] or "").strip()}
    return personal


def _fin_turno(fecha, entrada, salida):
    fin = datetime.combine(fecha, salida)
    ini = datetime.combine(fecha, entrada)
    if fin <= ini:
        fin += timedelta(days=1)
    return fin


def _calcular_filas(filas_lo, legajo: int | None) -> list[tuple[object, ResultadoDia]]:
    resultados = []
    salida_anterior = None
    for f in filas_lo:
        entrada_dt = None
        if f.entrada and not f.feriado:
            entrada_dt = datetime.combine(f.fecha, f.entrada)
        es_transporte = "transporte" in f.turno.lower()
        res = calcular_dia(
            f.fecha,
            f.entrada,
            f.salida,
            citacion=f.citacion,
            franco=not f.entrada or not f.salida,
            es_feriado=f.feriado,
            turno_transporte=es_transporte,
            salida_anterior=salida_anterior,
            entrada_actual_dt=entrada_dt,
        )
        if f.entrada and f.salida:
            salida_anterior = _fin_turno(f.fecha, f.entrada, f.salida)
        resultados.append((f, res))
    return resultados


def _diff_style(cell, diff: int):
    if diff == 0:
        return
    cell.fill = PatternFill("solid", fgColor="FFC7CE" if diff else "C6EFCE")
    cell.font = Font(color="9C0006" if diff else "006100")


def exportar_comparacion(
    nombre: str,
    legajo: int | None,
    pares: list[tuple[object, ResultadoDia]],
    destino: Path,
):
    wb = Workbook()
    ws = wb.active
    ws.title = "Comparacion"

    titulo = f"Comparación vs Legajos Online — {nombre}"
    if legajo:
        titulo += f" (Legajo {legajo})"
    ws["A1"] = titulo
    ws["A1"].font = Font(bold=True, size=13)
    ws.merge_cells("A1:R1")

    headers = [
        "Fecha", "Día", "Citación", "Entrada", "Salida", "Turno",
        "Total LO", "Total Calc", "Δ",
        "Norm LO", "Norm Calc", "Δ",
        "50% LO", "50% Calc", "Δ",
        "100% LO", "100% Calc", "Δ",
        "Noct LO", "Noct Calc", "Δ",
        "Cenas LO", "Cenas Calc",
    ]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=3, column=col, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F5496")
        c.alignment = Alignment(horizontal="center")

    row = 4
    diffs = 0
    for lo, calc in pares:
        campos = [
            (lo.fecha.strftime("%d-%m-%Y"), ""),
            (lo.dia, ""),
            (str(lo.citacion or "")[:8], ""),
            (str(lo.entrada or "")[:8], ""),
            (str(lo.salida or "")[:8], ""),
            (lo.turno[:12], ""),
            (lo.total_min, calc.total_min),
            (lo.horas_min, calc.horas_min),
            (lo.horas_50_min, calc.horas_50_min),
            (lo.horas_100_min, calc.horas_100_min),
            (lo.horas_nocturnas_min, calc.horas_nocturnas_min),
            (lo.cenas, calc.cenas),
        ]
        vals = [
            lo.fecha.strftime("%d-%m-%Y"), lo.dia,
            str(lo.citacion or "")[11:19] if lo.citacion else "",
            str(lo.entrada or "")[11:19] if lo.entrada else "",
            str(lo.salida or "")[11:19] if lo.salida else "",
            lo.turno[:14],
            _fmt_min(lo.total_min), _fmt_min(calc.total_min), calc.total_min - lo.total_min,
            _fmt_min(lo.horas_min), _fmt_min(calc.horas_min), calc.horas_min - lo.horas_min,
            _fmt_min(lo.horas_50_min), _fmt_min(calc.horas_50_min), calc.horas_50_min - lo.horas_50_min,
            _fmt_min(lo.horas_100_min), _fmt_min(calc.horas_100_min), calc.horas_100_min - lo.horas_100_min,
            _fmt_min(lo.horas_nocturnas_min), _fmt_min(calc.horas_nocturnas_min),
            calc.horas_nocturnas_min - lo.horas_nocturnas_min,
            lo.cenas, calc.cenas,
        ]
        for col, v in enumerate(vals, 1):
            ws.cell(row=row, column=col, value=v)
        for dcol in (9, 12, 15, 18, 21):
            d = vals[dcol - 1]
            if isinstance(d, int) and d != 0:
                _diff_style(ws.cell(row=row, column=dcol), d)
                diffs += 1
        row += 1

    ws.cell(row=row + 1, column=1, value=f"Diferencias marcadas en filas: {diffs}")
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    return diffs


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("archivo", type=Path, help="Excel exportado de Legajos Online")
    parser.add_argument("--legajo", type=int, default=None)
    args = parser.parse_args()

    if not args.archivo.exists():
        print(f"No existe: {args.archivo}")
        sys.exit(1)

    nombre, filas = importar_legajos_online(args.archivo)
    personal = _leer_personal()
    legajo = args.legajo or legajo_desde_nombre(nombre, personal)

    pares = _calcular_filas(filas, legajo)
    safe = nombre.replace(",", "").replace(" ", "_")[:30] or args.archivo.stem
    out = SALIDA / f"COMPARAR_{safe}.xlsx"
    diffs = exportar_comparacion(nombre, legajo, pares, out)

    print(f"Colaborador: {nombre}")
    print(f"Legajo detectado: {legajo}")
    print(f"Filas comparadas: {len(pares)}")
    print(f"Generado: {out}")

    # Resumen consola
    ok = sum(
        1 for lo, c in pares
        if lo.total_min == c.total_min and lo.horas_50_min == c.horas_50_min
        and lo.horas_100_min == c.horas_100_min
    )
    print(f"Coinciden total+50%+100%: {ok}/{len(pares)} días")


if __name__ == "__main__":
    main()
