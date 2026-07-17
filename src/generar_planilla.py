"""
Lee plantilla.xlsx, calcula horas y genera planillas por colaborador.
"""
from __future__ import annotations

import argparse
import sys
from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from calculo_horas import (
    ResultadoDia,
    _es_feriado,
    _es_franco,
    _es_transporte,
    _parse_date,
    _parse_time,
    calcular_dia,
    rango_periodo,
    sumar_minutos,
    _fmt_min,
)

BASE = Path(__file__).resolve().parent.parent
PLANTILLA = BASE / "data" / "plantilla.xlsx"
SALIDA_DIR = BASE / "data" / "planillas"


def _leer_personal(wb) -> dict[int, dict]:
    ws = wb["Personal"]
    personal = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        legajo = int(row[0])
        activo = str(row[2] or "si").strip().lower() in ("si", "sí", "s", "1", "true")
        if activo:
            personal[legajo] = {
                "legajo": legajo,
                "nombre": str(row[1] or "").strip(),
                "dni": str(row[3] or "").strip() if len(row) > 3 else "",
            }
    return personal


def _leer_asistencia(wb) -> list[dict]:
    ws = wb["Asistencia"]
    filas = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        fecha = _parse_date(row[0])
        if not fecha:
            continue
        legajo = int(row[1])
        filas.append({
            "fecha": fecha,
            "legajo": legajo,
            "entrada": _parse_time(row[2]),
            "salida": _parse_time(row[3]),
            "citacion": _parse_time(row[4]) if len(row) > 4 else None,
            "franco": _es_franco(row[5]) if len(row) > 5 else False,
            "feriado": _es_feriado(row[6]) if len(row) > 6 else False,
            "observacion": str(row[7] or "").strip() if len(row) > 7 else "",
            "turno_transporte": _es_transporte(row[8]) if len(row) > 8 else False,
        })
    return filas


def _salida_datetime(fecha: date, salida: time) -> datetime:
    dt = datetime.combine(fecha, salida)
    if salida < time(12, 0):
        pass
    return dt


def _fin_turno(fecha: date, entrada: time, salida: time) -> datetime:
    fin = datetime.combine(fecha, salida)
    ini = datetime.combine(fecha, entrada)
    if fin <= ini:
        fin += timedelta(days=1)
    return fin


def calcular_colaborador(
    legajo: int,
    personal: dict,
    asistencia: list[dict],
    anio: int,
    mes_fin: int,
) -> list[ResultadoDia]:
    dias_periodo = rango_periodo(anio, mes_fin)
    por_fecha = {a["fecha"]: a for a in asistencia if a["legajo"] == legajo}

    resultados: list[ResultadoDia] = []
    salida_anterior: datetime | None = None

    for fecha in dias_periodo:
        reg = por_fecha.get(fecha, {})
        entrada = reg.get("entrada")
        salida = reg.get("salida")
        franco = reg.get("franco", False)
        citacion = reg.get("citacion")
        feriado = reg.get("feriado", False)
        obs = reg.get("observacion", "")

        entrada_dt = None
        if entrada and not franco:
            entrada_dt = datetime.combine(fecha, entrada)

        res = calcular_dia(
            fecha,
            entrada,
            salida,
            citacion=citacion,
            franco=franco,
            es_feriado=feriado,
            turno_transporte=reg.get("turno_transporte", False),
            salida_anterior=salida_anterior,
            entrada_actual_dt=entrada_dt,
            observacion=obs,
        )
        resultados.append(res)

        if entrada and salida and not franco:
            salida_anterior = _fin_turno(fecha, entrada, salida)

    return resultados


def _estilo_encabezado(cell):
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="2F5496")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _borde():
    thin = Side(style="thin", color="AAAAAA")
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def exportar_planilla(
    legajo: int,
    info: dict,
    resultados: list[ResultadoDia],
    anio: int,
    mes_fin: int,
    destino: Path,
):
    wb = Workbook()
    ws = wb.active
    ws.title = f"Legajo {legajo}"

    titulo = f"Planilla de horas — {info['nombre']} — Legajo {legajo}"
    ws.merge_cells("A1:R1")
    ws["A1"] = titulo
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")

    periodo_txt = f"Período: 26/{mes_fin-1 if mes_fin > 1 else 12:02d}/{anio if mes_fin > 1 else anio-1} al 25/{mes_fin:02d}/{anio}"
    ws.merge_cells("A2:R2")
    ws["A2"] = periodo_txt
    ws["A2"].alignment = Alignment(horizontal="center")

    headers = [
        "Colaborador", "DNI", "Legajo", "Fecha", "Dia", "Feriado",
        "Citacion", "Entrada", "Salida", "Total Hs", "Horas",
        "Horas 50%", "Horas 100%", "Nocturnas", "Descanso", "Descanso 100%", "Cenas", "Dias",
    ]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=4, column=col, value=h)
        _estilo_encabezado(c)
        c.border = _borde()

    row = 5
    for r in resultados:
        vals = [
            info["nombre"],
            info.get("dni", ""),
            legajo,
            r.fecha.strftime("%d-%m-%Y"),
            r.dia_semana,
            "Si" if r.es_feriado else "No",
            r.citacion.strftime("%H:%M") if r.citacion else "",
            r.entrada.strftime("%H:%M") if r.entrada else ("Fco" if r.franco else ""),
            r.salida.strftime("%H:%M") if r.salida else ("Fco" if r.franco else ""),
            r.total_hs,
            r.horas,
            r.horas_50,
            r.horas_100,
            r.horas_nocturnas,
            r.descanso + (" !" if r.descanso_ok is False else ""),
            r.descanso_100,
            r.cenas,
            r.dias,
        ]
        for col, v in enumerate(vals, 1):
            cell = ws.cell(row=row, column=col, value=v)
            cell.border = _borde()
            if r.descanso_ok is False and col == 15:
                cell.fill = PatternFill("solid", fgColor="FFC7CE")
            if r.descanso_100_min > 0 and col == 16:
                cell.fill = PatternFill("solid", fgColor="FFEB9C")
        row += 1

    tot = sumar_minutos(resultados)
    tot_row = row + 1
    ws.cell(row=tot_row, column=9, value="TOTALES").font = Font(bold=True)
    ws.cell(row=tot_row, column=10, value=_fmt_min(tot["total"])).font = Font(bold=True)
    ws.cell(row=tot_row, column=11, value=_fmt_min(tot["horas"])).font = Font(bold=True)
    ws.cell(row=tot_row, column=12, value=_fmt_min(tot["horas_50"])).font = Font(bold=True)
    ws.cell(row=tot_row, column=13, value=_fmt_min(tot["horas_100"])).font = Font(bold=True)
    ws.cell(row=tot_row, column=14, value=_fmt_min(tot["nocturnas"])).font = Font(bold=True)
    ws.cell(row=tot_row, column=16, value=_fmt_min(tot["descanso_100"])).font = Font(bold=True)
    ws.cell(row=tot_row, column=17, value=tot["cenas"]).font = Font(bold=True)
    ws.cell(row=tot_row, column=18, value=tot["dias"]).font = Font(bold=True)

    for col in range(1, 19):
        ws.column_dimensions[get_column_letter(col)].width = 12
    ws.column_dimensions["A"].width = 28

    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)


def exportar_resumen(
    datos: list[tuple[int, dict, list[ResultadoDia]]],
    anio: int,
    mes_fin: int,
    destino: Path,
):
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    headers = ["Legajo", "Colaborador", "Total Hs", "Horas", "50%", "100%", "Cenas", "Dias", "Alertas descanso"]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        _estilo_encabezado(c)

    row = 2
    for legajo, info, resultados in datos:
        tot = sumar_minutos(resultados)
        alertas = sum(1 for r in resultados if r.descanso_ok is False)
        ws.append([
            legajo,
            info["nombre"],
            _fmt_min(tot["total"]),
            _fmt_min(tot["horas"]),
            _fmt_min(tot["horas_50"]),
            _fmt_min(tot["horas_100"]),
            tot["cenas"],
            tot["dias"],
            alertas,
        ])
        row += 1

    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)


def main():
    parser = argparse.ArgumentParser(description="Generar planillas de horas de transporte")
    parser.add_argument("--anio", type=int, default=2026)
    parser.add_argument("--mes", type=int, default=6, help="Mes de cierre del período (ej: 6 = 26/05 al 25/06)")
    parser.add_argument("--legajo", type=int, default=None, help="Solo un legajo")
    parser.add_argument("--plantilla", type=Path, default=PLANTILLA)
    args = parser.parse_args()

    if not args.plantilla.exists():
        print(f"No existe {args.plantilla}. Ejecutá primero: python src/crear_plantilla.py")
        sys.exit(1)

    wb = load_workbook(args.plantilla, data_only=True)
    personal = _leer_personal(wb)
    asistencia = _leer_asistencia(wb)

    legajos = [args.legajo] if args.legajo else sorted(personal.keys())
    datos = []
    for leg in legajos:
        if leg not in personal:
            print(f"Legajo {leg} no encontrado en Personal")
            continue
        res = calcular_colaborador(leg, personal, asistencia, args.anio, args.mes)
        info = personal[leg]
        datos.append((leg, info, res))
        nombre_safe = info["nombre"].replace(",", "").replace("/", "-")[:40]
        out = SALIDA_DIR / f"{leg}_{nombre_safe}_{args.mes:02d}-{args.anio}.xlsx"
        exportar_planilla(leg, info, res, args.anio, args.mes, out)
        print(f"Generada: {out}")

    if len(datos) > 1:
        resumen = SALIDA_DIR / f"RESUMEN_{args.mes:02d}-{args.anio}.xlsx"
        exportar_resumen(datos, args.anio, args.mes, resumen)
        print(f"Resumen: {resumen}")


if __name__ == "__main__":
    main()
