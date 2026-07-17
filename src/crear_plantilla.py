"""
Crea plantilla.xlsx con hojas Personal y Asistencia + datos de ejemplo.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

BASE = Path(__file__).resolve().parent.parent
OUT = BASE / "data" / "plantilla.xlsx"

# Personal extraído de las fotos (transporte)
PERSONAL = [
    (2794, "ACOSTA, Juan", ""),
    (324, "ALVAREZ, Mariano", ""),
    (319, "MOREYRA, Leandro", ""),
    (2787, "CABRERA, Andrea", ""),
    (2575, "CASTRO, Lucas", ""),
    (3210, "CORREA, Javier", ""),
    (247, "DIAZ, Silvia", ""),
    (253, "DOMINGUEZ, Sergio", ""),
    (2572, "FERNANDEZ, Ivan", ""),
    (2514, "FLORES, Laura", ""),
    (3573, "GOMEZ, Pablo", ""),
    (2824, "GONZALEZ, Hernan", ""),
    (67, "IBARRA, Ana", ""),
    (2510, "JUAREZ, Martin", ""),
    (2782, "LEDESMA, Gustavo", ""),
    (271, "LUNA, Tomas", ""),
    (174, "MARTINEZ, Diego", ""),
    (1392, "MEDINA, Federico", ""),
    (2680, "MENDEZ, Ricardo", ""),
    (123, "NUÑEZ, Carlos", ""),
    (1173, "OJEDA, Nicolas", ""),
    (3066, "PAEZ, Matias", ""),
    (285, "PERALTA, Juan", ""),
    (2939, "PEREZ, Mariano", ""),
    (2854, "PONCE, Leandro", ""),
    (3038, "RAMIREZ, Andrea", ""),
    (2784, "RIOS, Lucas", ""),
    (2758, "ROJAS, Javier", ""),
    (2537, "ROMERO, Silvia", ""),
    (2698, "RUIZ, Sergio", ""),
    (2823, "SANCHEZ, Ivan", ""),
    (267, "SUAREZ, Laura", ""),
    (2786, "TORRES, Pablo", ""),
    (2972, "VEGA, Hernan", ""),
    (2574, "VERA, Ana", ""),
    (2527, "VILLALBA, Martin", ""),
    (147, "ZALAZAR, Gustavo", ""),
    (77, "AGUIRRE, Tomas", ""),
    (2525, "CARRIZO, Diego", ""),
    (2664, "FIGUEROA, Federico", ""),
    (3571, "ACOSTA, Ricardo Javier", ""),
    (197, "ALVAREZ, Carlos Ricardo", ""),
    (3572, "MOREYRA, Nicolas Carlos", ""),
    (3283, "CABRERA, Matias Pablo", ""),
    (89, "CASTRO, Juan Mariano", ""),
    (69, "CORREA, Mariano Gustavo", ""),
    (325, "DIAZ, Leandro Javier", ""),
    (87, "DOMINGUEZ, Andrea Ricardo", ""),
    (2727, "FERNANDEZ, Lucas Carlos", ""),
    (224, "FLORES, Javier Pablo", ""),
    (2838, "GOMEZ, Silvia Mariano", ""),
    (85, "GONZALEZ, Sergio Gustavo", ""),
    (26, "IBARRA, Ivan Javier", ""),
    (268, "JUAREZ, Laura Ricardo", ""),
    (226, "LEDESMA, Pablo Carlos", ""),
    (258, "LUNA, Hernan Pablo", ""),
    (2934, "MARTINEZ, Ana Mariano", ""),
]

# Ejemplo de asistencia (parcial)
ASISTENCIA_EJEMPLO = [
    ("26/05/2026", 319, "12:56", "23:02", "13:10", "", "", ""),
    ("27/05/2026", 319, "07:03", "13:39", "07:00", "", "", ""),
    ("28/05/2026", 319, "", "", "", "si", "", "Franco"),
    ("29/05/2026", 319, "12:55", "22:02", "13:10", "", "", ""),
    ("30/05/2026", 319, "12:55", "22:02", "13:10", "", "", ""),
    ("31/05/2026", 319, "12:33", "18:44", "13:10", "", "", ""),
    ("26/05/2026", 2698, "11:56", "00:12", "", "", "", ""),
    ("27/05/2026", 2698, "17:41", "06:06", "", "", "", ""),
    ("28/05/2026", 2698, "", "", "", "si", "", "Fco"),
    ("29/05/2026", 2698, "05:33", "18:00", "", "", "", ""),
    ("30/05/2026", 2698, "05:30", "12:08", "", "", "", ""),
    ("31/05/2026", 2698, "", "", "", "si", "", "Fco"),
]


def _header(ws, row, headers, fill_color="2F5496"):
    fill = PatternFill("solid", fgColor=fill_color)
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(horizontal="center")


def main():
    wb = Workbook()

    # --- Personal ---
    ws_p = wb.active
    ws_p.title = "Personal"
    _header(ws_p, 1, ["Legajo", "Apellido y Nombre", "Activo", "DNI (opcional)"])
    for i, (leg, nom, dni) in enumerate(PERSONAL, 2):
        ws_p.append([leg, nom, "Si", dni])

    dv_activo = DataValidation(type="list", formula1='"Si,No"', allow_blank=True)
    ws_p.add_data_validation(dv_activo)
    dv_activo.add(f"C2:C500")

    ws_p.column_dimensions["A"].width = 10
    ws_p.column_dimensions["B"].width = 38
    ws_p.column_dimensions["C"].width = 8
    ws_p.column_dimensions["D"].width = 14

    # --- Asistencia ---
    ws_a = wb.create_sheet("Asistencia")
    _header(ws_a, 1, [
        "Fecha (dd/mm/aaaa)", "Legajo", "Entrada (HH:MM)", "Salida (HH:MM)",
        "Citacion (opc)", "Franco (si)", "Feriado (si)", "Observacion",
        "Turno (transporte/cronograma)",
    ])
    for fila in ASISTENCIA_EJEMPLO:
        ws_a.append(list(fila))

    dv_franco = DataValidation(type="list", formula1='"si,no"', allow_blank=True)
    ws_a.add_data_validation(dv_franco)
    dv_franco.add("F2:F5000")

    dv_fer = DataValidation(type="list", formula1='"si,no"', allow_blank=True)
    ws_a.add_data_validation(dv_fer)
    dv_fer.add("G2:G5000")

    for col, w in zip("ABCDEFGH", [16, 10, 14, 14, 14, 10, 10, 30]):
        ws_a.column_dimensions[col].width = w

    # --- Instrucciones ---
    ws_i = wb.create_sheet("Instrucciones")
    instrucciones = [
        "PLANILLA DE HORAS — TRANSPORTE",
        "",
        "1. PERSONAL: actualizá legajos y nombres. Columna Activo = Si/No.",
        "2. ASISTENCIA: cargá cada día quién trabajó.",
        "   - Fecha, Legajo, Entrada y Salida (formato 24hs, ej: 14:30)",
        "   - Franco: escribí 'si' si no trabajó ese día",
        "   - Feriado: 'si' si el día es feriado (todo se paga al 100%)",
        "3. Guardá el archivo y ejecutá generar_planilla.bat (o el comando Python).",
        "",
        "REGLAS DE CÁLCULO:",
        "- Período: del 26 al 25 de cada mes",
        "- Lun-Vie: jornada 8hs, extra al 50%, nocturnas 21:00-06:00 al 100%",
        "- Sábado: jornada 4hs, extra antes 13:00 al 50%, después al 100%",
        "- Domingo y feriados: todo al 100%",
        "- Descanso: alerta si hay menos de 12hs entre salida y entrada del día siguiente",
        "- Cenas: 1 si el turno pasa por la franja de cena (desde 20:00)",
    ]
    for i, linea in enumerate(instrucciones, 1):
        ws_i.cell(row=i, column=1, value=linea)
    ws_i.column_dimensions["A"].width = 80

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"Plantilla creada: {OUT}")


if __name__ == "__main__":
    main()
